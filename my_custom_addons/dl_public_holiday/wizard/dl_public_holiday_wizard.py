# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
import io
from datetime import datetime
import openpyxl

class DlPublicHolidayWizard(models.TransientModel):
    _name = 'dl.public.holiday.wizard'
    _description = 'Wizard Import/Export Ngày nghỉ lễ'

    action_type = fields.Selection([
        ('export', 'Tải file mẫu (Export)'),
        ('import', 'Nhập dữ liệu (Import)')
    ], string='Hành động', default='import', required=True)
    
    file_data = fields.Binary(string='File Excel', attachment=False)
    file_name = fields.Char(string='Tên file')

    def action_execute(self):
        self.ensure_one()
        if self.action_type == 'export':
            return self._export_template()
        else:
            return self._import_data()

    def _export_template(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ngày nghỉ lễ'

        # Headers
        headers = ['Tên ngày lễ (*)', 'Từ ngày (DD/MM/YYYY)', 'Đến ngày (DD/MM/YYYY)', 'Ghi chú']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header

        # Some sample data
        ws.append(['Tết Dương Lịch', '01/01/2026', '01/01/2026', 'Nghỉ tết dương'])
        ws.append(['Giỗ Tổ Hùng Vương', '26/04/2026', '26/04/2026', 'Mùng 10 tháng 3 Âm lịch'])
        ws.append(['Giải phóng Miền Nam', '30/04/2026', '30/04/2026', ''])
        ws.append(['Quốc tế Lao động', '01/05/2026', '01/05/2026', ''])

        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 40

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_base64 = base64.b64encode(output.read())

        # Update wizard
        self.write({
            'file_data': file_base64,
            'file_name': 'TEMPLATE_NGAY_LE.xlsx'
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'dl.public.holiday.wizard',
            'view_mode': 'form',
            'res_id': self.id,
            'views': [(False, 'form')],
            'target': 'new',
        }

    def _import_data(self):
        if not self.file_data:
            raise UserError(_("Vui lòng đính kèm file Excel để import!"))

        try:
            file_content = base64.b64decode(self.file_data)
            wb = openpyxl.load_workbook(filename=io.BytesIO(file_content), data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError(_("File không hợp lệ hoặc bị lỗi: %s") % str(e))

        HolidayObj = self.env['dl.public.holiday']
        company_id = self.env.company.id
        create_vals = []
        errors = []

        # Read rows starting from row 2
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]: # Skip empty rows
                continue
                
            name = str(row[0]).strip()
            date_from_str = row[1]
            date_to_str = row[2]
            note = str(row[3]).strip() if row[3] else ''

            if not date_from_str:
                errors.append(f"Dòng {row_idx}: Thiếu 'Từ ngày'.")
                continue

            # Parse date_from
            date_from = self._parse_date(date_from_str)
            if not date_from:
                errors.append(f"Dòng {row_idx}: 'Từ ngày' sai định dạng. Yêu cầu DD/MM/YYYY.")
                continue

            # Parse date_to
            date_to = False
            if date_to_str:
                date_to = self._parse_date(date_to_str)
                if not date_to:
                    errors.append(f"Dòng {row_idx}: 'Đến ngày' sai định dạng. Yêu cầu DD/MM/YYYY.")
                    continue
            else:
                date_to = date_from

            if date_from > date_to:
                errors.append(f"Dòng {row_idx}: 'Từ ngày' không được lớn hơn 'Đến ngày'.")
                continue

            create_vals.append({
                'name': name,
                'date_from': date_from,
                'date_to': date_to,
                'note': note,
                'company_id': company_id,
            })

        if errors:
            error_msg = "\n".join(errors)
            raise UserError(_("Phát hiện lỗi dữ liệu trong file:\n%s") % error_msg)

        if create_vals:
            HolidayObj.create(create_vals)

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Thành công'),
                'message': _('Đã import thành công %s ngày nghỉ lễ!') % len(create_vals),
                'sticky': False,
                'type': 'success',
                'next': {'type': 'ir.actions.act_window_close'}
            }
        }

    def _parse_date(self, value):
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if hasattr(value, 'date'):
            return value.date()
        try:
            return datetime.strptime(str(value).strip()[:10], '%d/%m/%Y').date()
        except ValueError:
            try:
                # Fallback to YYYY-MM-DD
                return datetime.strptime(str(value).strip()[:10], '%Y-%m-%d').date()
            except ValueError:
                return False
