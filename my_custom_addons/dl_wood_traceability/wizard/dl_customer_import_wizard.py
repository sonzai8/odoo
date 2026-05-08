# -*- coding: utf-8 -*-
import base64
import io
import openpyxl
from odoo import models, fields, api, _
from odoo.exceptions import UserError

class DlWoodCustomerImportWizard(models.TransientModel):
    _name = 'dl.wood.customer.import.wizard'
    _description = 'Wizard nhập/xuất khách hàng'

    file_data = fields.Binary(string='Chọn file Excel')
    file_name = fields.Char(string='Tên file')

    def action_export_template(self):
        """Xuất file Excel mẫu chứa danh sách khách hàng hiện có"""
        output = io.BytesIO()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Khách hàng"

        # Header
        headers = ['ID', 'Tên khách hàng', 'Mã số thuế', 'Điện thoại', 'Địa chỉ', 'Số lượng hợp đồng mẫu']
        for col_num, header in enumerate(headers, 1):
            sheet.cell(row=1, column=col_num).value = header

        # Data - chỉ lấy khách hàng
        partners = self.env['res.partner'].search([('customer_rank', '>', 0)])
        for row_num, partner in enumerate(partners, 2):
            sheet.cell(row=row_num, column=1).value = partner.id
            sheet.cell(row=row_num, column=2).value = partner.name
            sheet.cell(row=row_num, column=3).value = partner.vat or ''
            sheet.cell(row=row_num, column=4).value = partner.phone or ''
            sheet.cell(row=row_num, column=5).value = partner.contact_address or ''
            sheet.cell(row=row_num, column=6).value = len(partner.dl_contract_ids)

        workbook.save(output)
        file_content = base64.b64encode(output.getvalue())
        output.close()

        attachment = self.env['ir.attachment'].create({
            'name': 'Mau_Nhap_Khach_Hang.xlsx',
            'datas': file_content,
            'type': 'binary',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }

    def action_import_customers(self):
        """Nhập khách hàng từ file Excel"""
        if not self.file_data:
            raise UserError(_("Vui lòng chọn file Excel để nhập!"))

        try:
            file_content = base64.b64decode(self.file_data)
            workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            sheet = workbook.active
            
            count = 0
            # Giả sử header ở dòng 1, dữ liệu từ dòng 2
            for row in sheet.iter_rows(min_row=2, values_only=True):
                partner_id = row[0]
                name = row[1]
                vat = str(row[2]) if row[2] else ''
                phone = str(row[3]) if row[3] else ''
                address = row[4]

                if not name:
                    continue

                vals = {
                    'name': name,
                    'vat': vat,
                    'phone': phone,
                    'street': address,
                    'customer_rank': 1,
                }

                if partner_id:
                    try:
                        partner = self.env['res.partner'].browse(int(partner_id))
                        if partner.exists():
                            partner.write(vals)
                            count += 1
                            continue
                    except:
                        pass
                
                # Tạo mới nếu không có ID hoặc ID không hợp lệ
                self.env['res.partner'].create(vals)
                count += 1

        except Exception as e:
            raise UserError(_("Lỗi khi đọc file Excel: %s") % str(e))

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Thành công'),
                'message': _('Đã xử lý %s khách hàng.') % count,
                'type': 'success',
                'sticky': False,
            }
        }

    def action_misa_import_placeholder(self):
        """Placeholder cho chức năng nhập từ Misa"""
        raise UserError(_("Chức năng này chưa hoạt động."))
