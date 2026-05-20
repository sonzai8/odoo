# -*- coding: utf-8 -*-
import os
import sys

# Thêm đường dẫn Odoo vào sys.path để chạy script độc lập
sys.path.append('/Users/sonzai/dev/odoo 19/odoo')

import odoo
from odoo import api, tools

# Cấu hình môi trường Odoo
tools.config.parse_config(['-c', 'odoo.conf'])
db_name = 'odoo_db'

print(f"Connecting to database: {db_name}...")
registry = odoo.modules.registry.Registry(db_name)

with registry.cursor() as cr:
    env = api.Environment(cr, odoo.SUPERUSER_ID, {})
    
    print("--- KHỞI TẠO BÁO CÁO TIÊU HAO HỒ SƠ GỖ ---")
    
    # 1. Tạo bản ghi wizard mới
    wizard = env['dl.wood.dossier.consumption.wizard'].create({
        'production_state': 'all',
    })
    print(f"Wizard created with ID: {wizard.id}")
    
    # 2. Gọi action xuất file Excel
    action = wizard.action_export_excel()
    print("Action export completed successfully!")
    
    # 3. Kiểm tra xem file đã được tạo chưa
    if wizard.file_data:
        print("Success: file_data field is populated!")
        print(f"File Name: {wizard.file_name}")
        
        # Giải mã và đọc cấu trúc file Excel để kiểm thử tính đúng đắn
        import base64
        import openpyxl
        import io
        
        excel_bytes = base64.b64decode(wizard.file_data)
        workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        
        print("\nSheets found in generated Excel file:")
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            print(f" - Sheet: {sheet_name} (Rows: {sheet.max_row}, Cols: {sheet.max_column})")
            
            # Print first few rows to verify headers
            print("   Headers / First Row Data:")
            for r in range(1, min(6, sheet.max_row + 1)):
                row_vals = [sheet.cell(r, c).value for c in range(1, min(13, sheet.max_column + 1))]
                print(f"     Row {r}: {row_vals}")
                
        print("\n--- KIỂM THỬ THÀNH CÔNG RỰC RẼ! ---")
    else:
        print("Error: file_data field is empty!")
