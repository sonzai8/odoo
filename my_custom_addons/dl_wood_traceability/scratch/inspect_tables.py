import docx

def inspect():
    doc = docx.Document('my_custom_addons/dl_wood_traceability/static/TEMPLATES/TEMPLATE_HD_HSG.docx')
    print("Total tables:", len(doc.tables))
    for t_idx, table in enumerate(doc.tables):
        print(f"\n--- Table {t_idx} ---")
        for r_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"Row {r_idx}: {cells}")

if __name__ == '__main__':
    inspect()
