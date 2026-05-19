import os
import io
import docx
from docxtpl import DocxTemplate
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

def test_render_and_merge():
    # 1. Dummy context
    context = {
        'company_name': 'CÔNG TY TNHH ĐỨC LÂM',
        'company_address': 'Quảng Trị',
        'owner_name': 'Nguyễn Văn A',
        'date_of_hd': 'ngày 19 tháng 05 năm 2026',
        'total_quantity': '120',
        'price_total': '0',
        'main_species_name': 'Keo Lai, Bạch Đàn',
        'main_unit_name': 'm³',
        'main_quantity_val': '120',
        'table_rows': [
            {
                'stt': 1, 'idx': 1, 'species_name': 'Keo Lai', 'unit': 'm3', 'quantity': '80',
                'price_unit': 'Theo thỏa thuận', 'price_subtotal': '0', 'note': ''
            },
            {
                'stt': 2, 'idx': 2, 'species_name': 'Bạch Đàn', 'unit': 'm3', 'quantity': '40',
                'price_unit': 'Theo thỏa thuận', 'price_subtotal': '0', 'note': ''
            }
        ]
    }

    # 2. Render bằng docxtpl
    template_path = 'my_custom_addons/dl_wood_traceability/static/TEMPLATES/TEMPLATE_HD_HSG.docx'
    doc = DocxTemplate(template_path)
    doc.render(context)

    # 3. Post-process: Gộp ô cột đơn giá
    print("Post-processing merge...")
    for t_idx, table in enumerate(doc.tables):
        # Kiểm tra xem bảng có cột chứa "Theo thỏa thuận, theo giá thị trường" không
        target_col_idx = -1
        matching_rows = []
        
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                txt = cell.text.strip()
                if "Theo thỏa thuận, theo giá thị trường" in txt:
                    target_col_idx = c_idx
                    matching_rows.append(r_idx)
                    break # Chỉ cần tìm 1 cell chứa text này ở mỗi hàng
                    
        if target_col_idx != -1 and len(matching_rows) > 1:
            print(f"Found target table (Index {t_idx}), Column {target_col_idx}, Rows: {matching_rows}")
            
            # Thực hiện merge các ô
            start_row = matching_rows[0]
            end_row = matching_rows[-1]
            start_cell = table.cell(start_row, target_col_idx)
            end_cell = table.cell(end_row, target_col_idx)
            
            # Lưu lại văn bản gộp chung
            merged_text = "Theo thỏa thuận, theo giá thị trường tại từng thời điểm mua bán"
            
            # Merge
            start_cell.merge(end_cell)
            
            # Style ô sau khi gộp
            start_cell.text = ""
            p = start_cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(merged_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            
            # Căn giữa dọc
            start_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            print("Merged successfully!")

    # 4. Save
    output_path = 'my_custom_addons/dl_wood_traceability/scratch/output_merged.docx'
    doc.save(output_path)
    print(f"Saved to {output_path}")

if __name__ == '__main__':
    test_render_and_merge()
