# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
import base64
import io
import xlrd # or openpyxl

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

class SalaryKpiImportAttendanceWizard(models.TransientModel):
    _name = 'dl.salary.kpi.import.attendance.wizard'
    _description = 'Nhập chấm công KPI từ Excel'

    month_id = fields.Many2one('dl.salary.kpi.month', string='Tháng lương', required=True)
    file_data = fields.Binary(string='File Excel', required=True)
    file_name = fields.Char(string='Tên file')

    def action_import(self):
        if not self.file_data:
            return
        
        # Logic to read Excel horizontally
        # Column 1: Employee Name
        # Columns 2-32: Days 1-31
        
        file_content = base64.b64decode(self.file_data)
        wb = load_workbook(io.BytesIO(file_content), data_only=True)
        ws = wb.active
        
        att_types = self.env['dl.salary.kpi.attendance.type'].search([])
        att_type_map = {t.code: t.id for t in att_types}
        
        # Start from row 2 (assuming header is row 1)
        for row in range(2, ws.max_row + 1):
            emp_name = ws.cell(row=row, column=1).value
            if not emp_name:
                continue
            
            # Find the line in our system
            line = self.month_id.line_ids.filtered(lambda l: l.employee_name == emp_name)
            if not line:
                continue
            
            line = line[0]
            
            # Clear existing attendance for this line? 
            # Or just update? Usually safer to clear or update specifically.
            for day in range(1, 32):
                val = ws.cell(row=row, column=day + 1).value
                if val:
                    val = str(val).strip()
                    if val in att_type_map:
                        # Update or create attendance line
                        att_line = line.attendance_ids.filtered(lambda a: a.day == day)
                        if att_line:
                            att_line.write({'attendance_type_id': att_type_map[val]})
                        else:
                            self.env['dl.salary.kpi.attendance.line'].create({
                                'month_line_id': line.id,
                                'day': day,
                                'attendance_type_id': att_type_map[val]
                            })
                    elif val == '':
                         # Clear if empty?
                         att_line = line.attendance_ids.filtered(lambda a: a.day == day)
                         if att_line:
                             att_line.unlink()
        
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Thành công'),
                'message': _('Đã nhập dữ liệu chấm công từ Excel.'),
                'type': 'success',
            }
        }
