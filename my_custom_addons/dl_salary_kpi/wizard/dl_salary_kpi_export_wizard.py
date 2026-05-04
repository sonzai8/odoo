# -*- coding: utf-8 -*-
import base64
import io
import datetime
from odoo import models, fields, api, _
from odoo.tools import file_path
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import MergedCell
    from openpyxl.formula.translate import Translator
    import copy
except ImportError:
    load_workbook = None

class SalaryKpiExportWizard(models.TransientModel):
    _name = 'dl.salary.kpi.export.wizard'
    _description = 'Xuất bảng lương & KPI Excel (Skill Template 01)'

    month_id = fields.Many2one('dl.salary.kpi.month', string='Tháng lương', required=True)

    def _safe_write(self, ws, row, col, value):
        """Ghi dữ liệu an toàn vào ô, tránh ghi vào ô phụ của vùng gộp"""
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, MergedCell):
            # Tìm ô master của vùng gộp này
            for merged_range in ws.merged_cells.ranges:
                if cell.coordinate in merged_range:
                    master_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                    master_cell.value = value
                    return
        else:
            cell.value = value

    def _copy_row_formatting(self, ws, source_row, target_row):
        """Sao chép định dạng và dịch công thức từ dòng nguồn sang dòng đích"""
        for col in range(1, ws.max_column + 1):
            source_cell = ws.cell(row=source_row, column=col)
            target_cell = ws.cell(row=target_row, column=col)

            # Sao chép giá trị hoặc công thức
            if source_cell.data_type == 'f':  # Nếu là công thức
                # Dịch công thức từ dòng source_row sang target_row
                target_cell.value = Translator(source_cell.value, origin=source_cell.coordinate).translate_formula(target_cell.coordinate)
            else:
                target_cell.value = source_cell.value

            # Sao chép định dạng (Style)
            if source_cell.has_style:
                target_cell.font = copy.copy(source_cell.font)
                target_cell.border = copy.copy(source_cell.border)
                target_cell.fill = copy.copy(source_cell.fill)
                target_cell.number_format = copy.copy(source_cell.number_format)
                target_cell.protection = copy.copy(source_cell.protection)
                target_cell.alignment = copy.copy(source_cell.alignment)

    def action_export_xlsx(self):
        if not load_workbook:
            raise UserError(_("Thư viện openpyxl chưa được cài đặt."))

        # 1. Lấy template TEMPLATE_2026.xlsx
        try:
            template_path = file_path('dl_salary_kpi/static/src/templates/TEMPLATE_2026.xlsx')
        except FileNotFoundError:
            raise UserError(_("Không tìm thấy file mẫu Excel tại static/src/templates/TEMPLATE_2026.xlsx"))

        wb = load_workbook(template_path)
        ws = wb.active

        month_date = self.month_id.date_month
        if not month_date:
            raise UserError(_("Vui lòng chọn tháng/năm trên bảng cân đối."))

        # 2. Điền thông tin Header (Theo Mapping)
        # Ô A3: "Tháng MM năm YYYY"
        self._safe_write(ws, 3, 1, f"Tháng {month_date.strftime('%m')} năm {month_date.strftime('%Y')}")
        # Ô F4: Doanh thu công ty
        self._safe_write(ws, 4, 6, self.month_id.company_revenue)
        # Ô K4: Tháng
        self._safe_write(ws, 4, 11, int(month_date.strftime('%m')))
        # Ô O4: Năm
        self._safe_write(ws, 4, 15, int(month_date.strftime('%Y')))

        # 3. Điền dữ liệu nhân viên (Bắt đầu từ dòng 8)
        current_row = 8
        lines = self.month_id.line_ids
        
        for i, line in enumerate(lines):
            # Nếu không phải dòng đầu tiên (dòng 8), thực hiện copy định dạng từ dòng 8
            if current_row > 8:
                self._copy_row_formatting(ws, 8, current_row)
            
            # Mapping dữ liệu
            # Cột B: Họ và tên
            self._safe_write(ws, current_row, 2, line.employee_id.name)
            # Cột C: Mã số thuế
            self._safe_write(ws, current_row, 3, line.employee_id.dl_tax_id or '')
            # Cột D: Ngày sinh (dd/mm/yyyy)
            birthday_str = line.employee_id.birthday.strftime('%d/%m/%Y') if line.employee_id.birthday else ''
            self._safe_write(ws, current_row, 4, birthday_str)
            # Cột E: Giới tính (Nam/Nữ)
            gender = 'Nam' if line.employee_id.gender == 'male' else 'Nữ' if line.employee_id.gender == 'female' else ''
            self._safe_write(ws, current_row, 5, gender)
            # Cột F: Phòng ban Thuế
            self._safe_write(ws, current_row, 6, line.employee_id.dl_tax_department_id.name or '')
            # Cột G: Chức vụ thuế
            self._safe_write(ws, current_row, 7, line.employee_id.dl_tax_position or '')
            # Cột H: Lương cơ bản thuế
            self._safe_write(ws, current_row, 8, line.employee_id.dl_tax_base_salary or 0)

            # Chấm công (Cột I -> AM tương ứng 1 -> 31)
            for day in range(1, 32):
                col_idx = 8 + day # I là cột 9
                att_type = getattr(line, f'day_{day:02d}')
                self._safe_write(ws, current_row, col_idx, att_type.code if att_type else '')

            # Trợ cấp (Cột CP, CQ)
            # CP: Phụ cấp phụ nữ (Chỉ điền cho Nữ)
            if line.employee_id.gender == 'female':
                self._safe_write(ws, current_row, 94, line.payroll_women_allowance) # CP là cột 94
            else:
                self._safe_write(ws, current_row, 94, 0)
            
            # CQ: Ăn ca
            self._safe_write(ws, current_row, 95, line.payroll_meal_allowance) # CQ là cột 95

            # DG: Tiền KPI (Cân đối)
            if line.payroll_kpi_amount:
                self._safe_write(ws, current_row, 111, line.payroll_kpi_amount) # DG là cột 111

            # Số người phụ thuộc (Cột EI)
            total_dependents = len(line.employee_id.dependent_ids.filtered(lambda d: d.active))
            self._safe_write(ws, current_row, 139, total_dependents) # EI là cột 139

            current_row += 1

        # 4. Xuất file
        output = io.BytesIO()
        wb.save(output)
        file_data = base64.b64encode(output.getvalue())
        output.close()

        filename = f"BC_LUONG_KPI_{month_date.strftime('%m_%Y')}.xlsx"

        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': file_data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }
