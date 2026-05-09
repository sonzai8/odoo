# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
import io
import openpyxl
import time
import logging
from .. import constants
from datetime import datetime, date

_logger = logging.getLogger(__name__)

class SalaryKpiTaxImport(models.TransientModel):
    _name = 'dl.salary.kpi.tax.import'
    _description = 'Wizard nhập thông tin thuế nhân viên'

    file_data = fields.Binary(string='File Excel')
    file_name = fields.Char(string='Tên file')
    state = fields.Selection([('init', 'Chọn File'), ('preview', 'Xem trước')], default='init')
    preview_html = fields.Html(string='Xem trước dữ liệu')

    def action_export_template(self):
        """Xuất file mẫu chứa danh sách nhân viên hiện có"""
        return {
            'type': 'ir.actions.act_url',
            'url': '/dl_salary_kpi/export_tax_employees',
            'target': 'new',
        }

    def _parse_date(self, value):
        if not value:
            return False
        if isinstance(value, (date, datetime)):
            return value
        if isinstance(value, str):
            value = value.strip()
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y'):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError:
                    continue
        return False

    def _read_excel_data(self):
        if not self.file_data:
            raise UserError("Vui lòng chọn file Excel trước khi thực hiện.")
        
        file_content = base64.b64decode(self.file_data)
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        ws = wb.active

        # --- BƯỚC 1: KIỂM TRA ĐỊNH DẠNG TEMPLATE ---
        # Kiểm tra dòng 1 hoặc dòng 2 xem có phải là tiêu đề chuẩn không
        header_found = False
        start_row = 2
        
        # Thử kiểm tra 3 dòng đầu tiên để tìm header
        for r in range(1, 4):
            row_vals = [str(ws.cell(row=r, column=c).value or '').strip() for c in range(1, 15)]
            # Kiểm tra các cột then chốt: Họ và Tên (cột 3), MST (cột 10) - Không phân biệt hoa thường
            val_c3 = row_vals[2].lower()
            val_c10 = row_vals[9].lower()
            header_full_name = constants.COL_FULL_NAME.lower()
            header_tax_id = constants.COL_TAX_ID.lower()
            
            if "họ" in val_c3 and "tên" in val_c3 and "mã số thuế" in val_c10:
                header_found = True
                start_row = r + 1
                break
        
        if not header_found:
            raise UserError(_("Sai định dạng file Excel! Hệ thống không tìm thấy các cột '%s' và '%s' ở vị trí mong đợi.") % (constants.COL_FULL_NAME, constants.COL_TAX_ID))

        parsed_data = []
        # Bắt đầu đọc từ dòng sau tiêu đề
        for row_idx, row in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start_row):
            if not row or not any(row[1:3]): # Kiểm tra có dữ liệu cơ bản không (ID hoặc Tên)
                continue
            
            # Cấu trúc: 0:STT, 1:ID NV, 2:Họ tên, 3:Tên riêng, 4:CCCD, 5:Ngày sinh, 6:Email, 7:SĐT, 8:Giới tính, 9:MST, 10:Chức vụ, 11:Phòng ban, 12:Lương, 13:Ngày nghỉ
            
            # Kiểm tra và xử lý Lương (Cột 12) - Tránh lỗi nếu vấp phải dòng text hoặc rỗng
            raw_salary = row[12] if len(row) > 12 else 0.0
            try:
                base_salary = float(raw_salary) if raw_salary else 0.0
            except (ValueError, TypeError):
                # Nếu không phải số, có thể là dòng ghi chú hoặc header thừa, bỏ qua
                continue

            emp_id = int(row[1]) if row[1] and str(row[1]).isdigit() else False
            name = str(row[2]).strip() if row[2] else ''
            dl_tax_id = str(row[9]).strip() if len(row) > 9 and row[9] else ''
            dept_name = str(row[11]).strip() if len(row) > 11 and row[11] else ''
            
            error = ""
            if not name:
                error = _("Thiếu %s.") % constants.COL_FULL_NAME
            elif not dl_tax_id:
                error = "Thiếu Mã số thuế."
            
            employee = False
            if emp_id:
                employee = self.env['hr.employee'].browse(emp_id)
                if not employee.exists():
                    employee = False
            
            if not employee and name and dl_tax_id:
                employee = self.env['hr.employee'].search([('name', '=', name), ('dl_tax_id', '=', dl_tax_id)], limit=1)

            parsed_data.append({
                'row_idx': row_idx,
                'emp_id': emp_id,
                'name': name,
                'dl_first_name': str(row[3]) if len(row) > 3 and row[3] else '',
                'identification_id': str(row[4]) if len(row) > 4 and row[4] else '',
                'birthday': self._parse_date(row[5]) if len(row) > 5 else False,
                'email': str(row[6]) if len(row) > 6 and row[6] else '',
                'work_phone': str(row[7]) if len(row) > 7 and row[7] else '',
                'sex': str(row[8]) if len(row) > 8 and row[8] else '',
                'dl_tax_id': dl_tax_id,
                'dl_tax_position': str(row[10]) if len(row) > 10 and row[10] else '',
                'dl_tax_department_name': dept_name,
                'dl_tax_base_salary': base_salary,
                'dl_departure_date': self._parse_date(row[13]) if len(row) > 13 else False,
                'x_bank_account': str(row[14]).strip() if len(row) > 14 and row[14] else '',
                'x_bank_name': str(row[15]).strip() if len(row) > 15 and row[15] else '',
                'error': error
            })
        return parsed_data

    def action_preview(self):
        parsed_data = self._read_excel_data()
        if not parsed_data:
            raise UserError("File Excel không có dữ liệu hợp lệ.")

        valid_count = sum(1 for d in parsed_data if not d['error'])
        error_count = len(parsed_data) - valid_count

        preview_rows = parsed_data[:5]
        if len(parsed_data) > 10:
            preview_rows.extend([None])
            preview_rows.extend(parsed_data[-5:])
        elif len(parsed_data) > 5:
            preview_rows.extend(parsed_data[5:])

        html = f"""
        <div class="alert alert-info">
            <strong>Tổng số nhân viên trong file:</strong> {len(parsed_data)}<br/>
            <strong>Hợp lệ:</strong> <span class="text-success">{valid_count}</span> | 
            <strong>Lỗi:</strong> <span class="text-danger">{error_count}</span>
        </div>
        <table class="table table-sm table-bordered">
            <thead class="table-light">
                <tr>
                    <th>Dòng</th>
                    <th>{constants.COL_ID_NV}</th>
                    <th>{constants.COL_FULL_NAME}</th>
                    <th>{constants.COL_TAX_ID}</th>
                    <th>CCCD</th>
                    <th>{constants.COL_TAX_DEPARTMENT}</th>
                    <th>{constants.COL_TAX_BASE_SALARY}</th>
                    <th>Trạng thái</th>
                </tr>
            </thead>
            <tbody>
        """

        for row in preview_rows:
            if row is None:
                html += """<tr><td colspan="8" class="text-center"><strong>...</strong></td></tr>"""
                continue
            
            row_class = "table-danger" if row['error'] else ""
            status_text = f"<span class='text-danger'>{row['error']}</span>" if row['error'] else "<span class='text-success'>Hợp lệ</span>"
            
            html += f"""
                <tr class="{row_class}">
                    <td>{row['row_idx']}</td>
                    <td>{row['emp_id'] or ''}</td>
                    <td>{row['name']}</td>
                    <td>{row['dl_tax_id']}</td>
                    <td>{row['identification_id']}</td>
                    <td>{row['dl_tax_department_name']}</td>
                    <td>{row['dl_tax_base_salary']:,.0f}</td>
                    <td>{status_text}</td>
                </tr>
            """
        html += "</tbody></table>"

        self.preview_html = html
        self.state = 'preview'
        
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'dl.salary.kpi.tax.import',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_back(self):
        self.state = 'init'
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'dl.salary.kpi.tax.import',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_confirm_import(self):
        t0 = time.time()
        parsed_data = self._read_excel_data()
        t1 = time.time()
        _logger.info("=== IMPORT LOG: Đọc và parse Excel mất: %.3fs ===", t1 - t0)

        if not parsed_data:
            return
        
        count_create = 0
        count_update = 0
        
        # 1. Tối ưu tìm kiếm: Thu thập MST và Tên để search 1 lần
        tax_ids = [r['dl_tax_id'] for r in parsed_data if r['dl_tax_id']]
        names = [r['name'] for r in parsed_data if r['name']]
        
        # Tạo mapping để tìm nhanh nhân viên hiện có
        existing_emps_by_id = {}
        emp_ids_to_browse = [r['emp_id'] for r in parsed_data if r['emp_id']]
        if emp_ids_to_browse:
            emps = self.env['hr.employee'].browse(emp_ids_to_browse).exists()
            existing_emps_by_id = {e.id: e for e in emps}
            
        existing_emps_by_key = {}
        domain = ['|', ('dl_tax_id', 'in', tax_ids), ('name', 'in', names)]
        all_emps = self.env['hr.employee'].search(domain)
        for e in all_emps:
            key = (e.name, e.dl_tax_id)
            if key not in existing_emps_by_key:
                existing_emps_by_key[key] = e
        
        t2 = time.time()
        _logger.info("=== IMPORT LOG: Tìm kiếm và tạo Mapping mất: %.3fs ===", t2 - t1)

        # 2. Tối ưu ghi dữ liệu: Tắt tracking và mail để tăng tốc
        optimized_context = {
            'tracking_disable': True, 
            'mail_notrack': True,
            'no_reset_password': True,
            'prefetch_fields': False,
            'recompute': False # Tạm dừng tính toán lại trong vòng lặp nếu có thể
        }
        Employee = self.env['hr.employee'].with_context(**optimized_context)
        
        # Cache phòng ban thuế
        dept_cache = {d.name: d.id for d in self.env['dl.tax.department'].search([])}

        for row in parsed_data:
            if row['error']:
                continue

            dept_id = dept_cache.get(row['dl_tax_department_name'])
            vals = {
                'name': row['name'],
                'birthday': row['birthday'],
                'dl_first_name': row['dl_first_name'],
                'email': row['email'],
                'work_phone': row['work_phone'],
                'identification_id': row['identification_id'],
                'sex': 'male' if str(row['sex'] or '').strip().upper() == 'NAM' else 'female' if str(row['sex'] or '').strip().upper() == 'NỮ' else 'other',
                'dl_tax_id': row['dl_tax_id'],
                'dl_tax_position': row['dl_tax_position'],
                'dl_tax_department_id': dept_id,
                'dl_tax_base_salary': row['dl_tax_base_salary'],
                'dl_departure_date': row['dl_departure_date'],
                'x_bank_account': row['x_bank_account'],
                'x_bank_name': row['x_bank_name'],
            }
            
            employee = existing_emps_by_id.get(row['emp_id'])
            if not employee:
                employee = existing_emps_by_key.get((row['name'], row['dl_tax_id']))

            if employee:
                # ÉP CONTEXT TỐI ƯU CHO LỆNH WRITE
                # Chỉ write nếu thực sự có dữ liệu thay đổi để tiết kiệm CPU
                changed_vals = {}
                for field, value in vals.items():
                    old_val = getattr(employee, field)
                    
                    # Xử lý so sánh cho trường Many2one (Phòng ban)
                    if field.endswith('_id') and hasattr(old_val, 'id') and not isinstance(old_val, str):
                        old_val = old_val.id
                    
                    # Xử lý so sánh cho trường Ngày tháng
                    if isinstance(old_val, date) and isinstance(value, str):
                        value = self._parse_date(value)
                    
                    if old_val != value:
                        changed_vals[field] = value
                
                if changed_vals:
                    employee.with_context(**optimized_context).write(changed_vals)
                    count_update += 1
            else:
                new_emp = Employee.create(vals)
                existing_emps_by_key[(new_emp.name, new_emp.dl_tax_id)] = new_emp
                count_create += 1
        
        t3 = time.time()
        _logger.info("=== IMPORT LOG: Vòng lặp Ghi dữ liệu (Write/Create) mất: %.3fs ===", t3 - t2)
        _logger.info("=== TOTAL IMPORT TIME: %.3fs ===", t3 - t0)
            
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Thành công'),
                'message': _('Tốc độ xử lý đã được tối ưu. Đã xong: %s cập nhật, %s thêm mới.') % (count_update, count_create),
                'type': 'success',
                'sticky': False,
            }
        }
