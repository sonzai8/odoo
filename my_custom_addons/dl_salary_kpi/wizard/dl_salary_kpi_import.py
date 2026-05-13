# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError
import base64
import io
import openpyxl
from .. import constants

class SalaryKpiImportWizard(models.TransientModel):
    _name = 'dl.salary.kpi.import.wizard'
    _description = 'Wizard nhập bảng công'

    month_id = fields.Many2one('dl.salary.kpi.month', string='Tháng bảng công')
    wizard_type = fields.Selection([
        ('normal', 'Công Thường'),
        ('overtime', 'Làm Thêm'),
        ('internal_salary', 'Lương nội bộ (KPI Target)')
    ], string='Loại xử lý', default='normal')
    file_data = fields.Binary(string='File Excel', required=False)
    file_name = fields.Char(string='Tên file')

    def action_export(self):
        self.ensure_one()
        if self.wizard_type == 'overtime':
            return self.month_id.action_export_ot_excel()
        if self.wizard_type == 'internal_salary':
            return self.month_id.action_export_internal_salary_excel()
        return self.month_id.action_export_excel()

    def action_import(self):
        if not self.file_data:
            return
        
        file_content = base64.b64decode(self.file_data)
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        if 'Bang Cham Cong' in wb.sheetnames:
            ws = wb['Bang Cham Cong']
        else:
            ws = wb.active # Fallback

        if self.wizard_type == 'internal_salary':
            return self._import_internal_salary(ws)

        # Lấy bản đồ mã công -> ID
        att_types = self.env['dl.salary.kpi.attendance.type'].search([])
        att_type_map = {t.code: t.id for t in att_types}
        att_code_map = {t.id: t.code for t in att_types}

        errors = []
        import_data = [] # List of (line_record, values_to_write)

        # Duyệt từ dòng 4 (skip headers)
        count_skipped_departure = 0
        from datetime import date

        for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), 4):

            tax_id_excel = str(row[1]).strip() if row[1] else ""
            if tax_id_excel.endswith('.0'):
                tax_id_excel = tax_id_excel[:-2]
                
            emp_name_excel = str(row[2]).strip() if row[2] else ""
            
            if not emp_name_excel:
                continue


            # Tìm line tương ứng trong tháng dựa trên Tên và MST
            line = self.month_id.line_ids.filtered(
                lambda l: (l.employee_id.name or '').strip() == emp_name_excel and 
                          (l.employee_id.dl_tax_id or '').strip() == tax_id_excel
            )
            
            if not line:
                errors.append(f"Dòng {row_idx}: Không tìm thấy nhân viên '{emp_name_excel}' có MST '{tax_id_excel}' trong bảng công tháng này.")
                continue
            
            if len(line) > 1:
                errors.append(f"Dòng {row_idx}: Tìm thấy {len(line)} nhân viên trùng Tên và MST ({emp_name_excel} - {tax_id_excel}). Vui lòng kiểm tra lại dữ liệu.")
                continue

            # Kiểm tra ngày nghỉ việc
            departure_date = fields.Date.to_date(line.employee_id.dl_departure_date)
            month_date = self.month_id.date_month
            year, month = month_date.year, month_date.month

            vals = {}
            row_codes = {} # To check shift change
            has_departure_skip = False
            
            # Lấy dữ liệu hiện tại từ hệ thống để so sánh chuỗi
            for day in range(1, 32):
                field_name = f'day_{day:02d}' if self.wizard_type == 'normal' else f'ot_day_{day:02d}'
                current_att = getattr(line, field_name)
                row_codes[day] = current_att.code if current_att else False

            # Ghi đè bằng dữ liệu từ Excel
            # Công thường: Cột 8 (index 7). Làm thêm: Cột 44 (index 43)
            col_offset = 7 if self.wizard_type == 'normal' else 43
            for day in range(1, 32):
                field_name = f'day_{day:02d}' if self.wizard_type == 'normal' else f'ot_day_{day:02d}'
                
                # Check departure date
                try:
                    d = date(year, month, day)
                    # Kiểm tra ngày nghỉ việc: Nếu ngày đang xét >= ngày nghỉ việc
                    if departure_date and d >= departure_date:
                        if self.wizard_type == 'normal':
                            # Ép buộc dùng mã NV cho công thường cho TẤT CẢ các ngày từ khi nghỉ (trừ Chủ Nhật)
                            if d.weekday() < 6:
                                att_nv_id = att_type_map.get('NV')
                                vals[field_name] = att_nv_id
                                row_codes[day] = 'NV'
                            else:
                                vals[field_name] = False
                                row_codes[day] = False
                        else:
                            # Đối với công làm thêm, sau khi nghỉ việc thì không có công LT
                            vals[field_name] = False
                            row_codes[day] = False
                        
                        has_departure_skip = True
                        # Quan trọng: continue ngay tại đây để không đọc dữ liệu từ Excel cho ngày này
                        continue
                    
                    # Quy tắc: Chấm công làm thêm chỉ cho phép vào ngày Chủ Nhật (bỏ qua check này khi import theo yêu cầu mới: cho phép import toàn bộ)
                    # if self.wizard_type != 'normal' and d.weekday() != 6:
                    #     continue 
                    pass
                except ValueError:
                    pass

                # Excel index: col_offset + day - 1
                try:
                    val = row[col_offset + day - 1]
                    code = str(val).strip().upper() if val else ""
                except IndexError:
                    code = ""

                if code:
                    if code not in att_type_map:
                        errors.append(f"Dòng {row_idx}: Mã công '{code}' ngày {day:02d} không hợp lệ.")
                    else:
                        # Kiểm tra xem mã công có phải là mã làm thêm không (nếu đang import OT)
                        if self.wizard_type == 'overtime':
                            att_type_id = att_type_map[code]
                            # Ta cần lấy thông tin apply_to của mã này
                            # Để tối ưu, ta có thể đã cache hoặc search lại.
                            # Vì số lượng dòng ít, search lại hoặc dùng env cache
                            att_type = self.env['dl.salary.kpi.attendance.type'].browse(att_type_id)
                            if att_type.apply_to == 'normal':
                                errors.append(f"Dòng {row_idx}: Mã '{code}' ngày {day:02d} không phải là mã chấm công làm thêm.")
                            else:
                                vals[field_name] = att_type_id
                                row_codes[day] = code
                        else:
                            vals[field_name] = att_type_map[code]
                            row_codes[day] = code
                else:
                    vals[field_name] = False
                    row_codes[day] = False

            if has_departure_skip:
                count_skipped_departure += 1

            # Kiểm tra quy tắc đổi ca (chỉ cho công thường)
            if self.wizard_type == 'normal':
                for i in range(1, 31):
                    cur = row_codes.get(i)
                    nxt = row_codes.get(i+1)
                    if cur == 'Đ' and nxt == 'N':
                        errors.append(f"Dòng {row_idx} ({emp_name_excel}): Lỗi đổi ca Đ sang N tại ngày {i:02d}-{i+1:02d} (Thiếu ĐC).")

            if vals and not errors:
                import_data.append((line, vals))

        if errors:
            # Chỉ hiển thị tối đa 10 lỗi đầu tiên
            display_errors = errors[:10]
            if len(errors) > 10:
                display_errors.append(f"... và còn {len(errors) - 10} lỗi khác nữa.")
            raise UserError("\n".join(display_errors))

        # Nếu không có lỗi nào thì mới tiến hành lưu
        for line, vals in import_data:
            line.write(vals)

        msg = _('Đã cập nhật dữ liệu cho %s nhân viên.') % len(import_data)
        if count_skipped_departure > 0:
            msg += _('\nLưu ý: Có %s nhân viên bị bỏ qua các ngày sau ngày nghỉ việc.') % count_skipped_departure

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Thành công'),
                'message': msg,
                'type': 'success',
                'sticky': True if count_skipped_departure > 0 else False,
            }
        }

    def _import_internal_salary(self, ws):
        """Import Lương nội bộ (Ln) dựa trên CCCD (Cột B)"""
        errors = []
        import_data = []
        
        # Duyệt từ dòng 4
        for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), 4):
            cccd = str(row[1]).strip() if row[1] else False
            if not cccd:
                continue
            
            # Làm sạch chuỗi CCCD (nếu là số thì bỏ .0)
            if cccd.endswith('.0'):
                cccd = cccd[:-2]

            # Validate độ dài CCCD chuẩn (9 số CMND cũ hoặc 12 số CCCD mới)
            if len(cccd) not in [9, 12]:
                errors.append(f"Dòng {row_idx}: Số CCCD '{cccd}' không đúng định dạng (phải là 9 hoặc 12 số).")
                continue
                
            emp_name_excel = str(row[2]).strip() if row[2] else ""
            
            # Tìm line tương ứng dựa trên identification_id (Trim cả hai đầu)
            line = self.month_id.line_ids.filtered(lambda l: (l.identification_id or '').strip() == cccd)
            if not line:
                errors.append(f"Dòng {row_idx}: Không tìm thấy nhân viên có CCCD '{cccd}' trong bảng công tháng này.")
                continue
            
            if len(line) > 1:
                errors.append(f"Dòng {row_idx}: Tìm thấy nhiều dòng có cùng CCCD '{cccd}'.")
                continue
                
            # Lương nội bộ ở cột F (index 5)
            try:
                internal_salary = float(row[5]) if row[5] else 0.0
            except (ValueError, TypeError):
                errors.append(f"Dòng {row_idx}: Lương nội bộ '{row[5]}' không hợp lệ.")
                continue
                
            import_data.append((line, {'payroll_internal_salary': internal_salary}))

        if errors:
            display_errors = errors[:10]
            if len(errors) > 10:
                display_errors.append(f"... và còn {len(errors) - 10} lỗi khác nữa.")
            raise UserError("\n".join(display_errors))

        for line, vals in import_data:
            line.write(vals)

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Thành công'),
                'message': _('Đã cập nhật Lương nội bộ cho %s nhân viên.') % len(import_data),
                'type': 'success',
            }
        }

    def action_import_internal_placeholder(self):
        """Placeholder cho chức năng import công từ file nội bộ"""
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Đang Phát Triển',
                'message': 'Đây là chức năng sẽ phát triển thêm. Nhập công từ file chấm công nội bộ. Không cần phải tốn thêm 1 bước chuẩn hoá công từ nội bộ ra công bên ngoài nữa.',
                'type': 'warning',
                'sticky': True,
            }
        }
