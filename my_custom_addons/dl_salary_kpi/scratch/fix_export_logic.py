# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.tools import file_path
import logging
import threading
import io
import base64
from datetime import date
from calendar import monthrange
import time

_logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
except ImportError:
    load_workbook = None
    PatternFill = None

class SalaryKpiMonth(models.Model):
    _inherit = 'dl.salary.kpi.month'

    def _bg_generate_excel_report(self, attachment_id, report_type):
        """
        Hàm chạy ngầm (Background Thread) để tạo file Excel.
        Sử dụng new cursor để tránh xung đột transaction.
        """
        new_cr = self.pool.cursor()
        self = self.with_env(self.env(cr=new_cr))
        try:
            if report_type == 'salary':
                attachment = self.env['ir.attachment'].browse(attachment_id)
                result = self._get_salary_report_data()
                attachment.write({'datas': result['datas']})
            elif report_type == 'kpi':
                attachment = self.env['ir.attachment'].browse(attachment_id)
                from odoo.addons.dl_salary_kpi.models import kpi_export_logic
                file_data, filename = kpi_export_logic.export_kpi_point_excel(self)
                attachment.write({'datas': file_data})
            
            new_cr.commit()
        except Exception as e:
            _logger.error("Lỗi khi tạo file Excel trong background: %s", str(e))
            new_cr.rollback()
        finally:
            new_cr.close()

    def _get_salary_report_data(self):
        """Hàm nội bộ tách logic tạo dữ liệu Excel Lương."""
        if not load_workbook:
            raise UserError(_("Thư viện openpyxl chưa được cài đặt."))

        template_path = file_path('dl_salary_kpi/static/src/templates/TEMPLATE_2026.xlsx')
        wb = load_workbook(template_path)
        wb.calculation.fullCalcOnLoad = True
        ws = wb.active
        month_date = self.date_month
        ws.title = f"Tháng {month_date.strftime('%m')} - năm {month_date.strftime('%Y')}"

        self._safe_write(ws, 3, 1, f"Tháng {month_date.strftime('%m')} năm {month_date.strftime('%Y')}")
        self._safe_write(ws, 4, 7, self.dl_revenue)
        self._safe_write(ws, 4, 12, int(month_date.strftime('%m')))
        self._safe_write(ws, 4, 16, int(month_date.strftime('%Y')))
        
        bonus_names = [b.name for b in self.bonus_line_ids]
        self._safe_write(ws, 5, 133, " + ".join(bonus_names) if bonus_names else "")

        current_row = 8
        last_day = monthrange(month_date.year, month_date.month)[1]

        for i, line in enumerate(self.line_ids):
            if current_row > 8:
                self._copy_row_formatting(ws, 8, current_row)
            
            # Ghi thông tin nhân viên
            self._safe_write(ws, current_row, 1, i + 1)
            self._safe_write(ws, current_row, 2, line.employee_id.dl_tax_id or '')
            self._safe_write(ws, current_row, 3, line.employee_id.name)
            self._safe_write(ws, current_row, 4, line.employee_id.birthday.strftime('%d/%m/%Y') if line.employee_id.birthday else '')
            self._safe_write(ws, current_row, 5, line.employee_id.identification_id or '')
            gender = 'Nam' if line.employee_id.sex == 'male' else 'Nữ' if line.employee_id.sex == 'female' else ''
            self._safe_write(ws, current_row, 6, gender)
            self._safe_write(ws, current_row, 7, line.employee_id.dl_tax_department_id.name or '')
            self._safe_write(ws, current_row, 8, line.employee_id.dl_tax_position or '')
            self._safe_write(ws, current_row, 9, line.employee_id.dl_tax_base_salary or 0)

            # 1. Ghi mã công thực tế và công làm thêm
            for day in range(1, 32):
                col_idx = 9 + day
                att_type = getattr(line, f'day_{day:02d}')
                # Reset màu nền cho chắc chắn (vì hàng 8 có thể có format cũ)
                ws.cell(row=current_row, column=col_idx).fill = PatternFill(fill_type=None)
                self._safe_write(ws, current_row, col_idx, att_type.code if att_type else '')

                if day <= last_day:
                    ot_att = getattr(line, f'ot_day_{day:02d}')
                    self._safe_write(ws, current_row, 45 + day, ot_att.code if ot_att else '')

            # 2. Hậu xử lý: Tự động điền CP cho các ngày trống trong tuần (Chỉ khi đã Confirmed)
            if self.state == 'confirmed':
                for day in range(1, last_day + 1):
                    current_date = date(month_date.year, month_date.month, day)
                    col_idx = 9 + day
                    cell = ws.cell(row=current_row, column=col_idx)
                    # Nếu là T2-T7 và chưa có mã công
                    if current_date.weekday() < 6 and not cell.value:
                        cell.value = "CP"
                        cell.fill = PatternFill(start_color='FFFFE0', end_color='FFFFE0', fill_type='solid')

            # Ghi các thông tin phụ cấp và thưởng
            self._safe_write(ws, current_row, 95, self.dl_women_allowance if line.employee_id.sex == 'female' else 0)
            self._safe_write(ws, current_row, 96, self.dl_meal_allowance)
            self._safe_write(ws, current_row, 111, line.payroll_kpi_amount or 0)
            self._safe_write(ws, current_row, 135, line.payroll_annual_bonus or 0)
            self._safe_write(ws, current_row, 138, line.payroll_pit_number_of_dependents or 0)
            self._safe_write(ws, current_row, 120, line.payroll_deduction_tncn or 0)
            
            total_bonus = 0
            for bonus in self.bonus_line_ids:
                if bonus.gender == 'all' or bonus.gender == line.employee_id.sex:
                    total_bonus += bonus.amount
            self._safe_write(ws, current_row, 133, total_bonus)
            
            current_row += 1

        output = io.BytesIO()
        wb.save(output)
        file_data = base64.b64encode(output.getvalue())
        output.close()
        
        return {'datas': file_data, 'filename': f"BC_LUONG_KPI_{month_date.strftime('%m_%Y')}.xlsx"}
