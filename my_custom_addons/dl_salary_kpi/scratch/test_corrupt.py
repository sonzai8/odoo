import openpyxl
from copy import copy

wb = openpyxl.load_workbook('my_custom_addons/dl_salary_kpi/static/src/templates/TEMPLATE_KPI_2026.xlsx')
ws = wb['TEMPLATE_KPI_AMOUNT']

# simulate the writes
ws.cell(1,1).value = 'Company'
ws.cell(2,1).value = 'MST'
ws.cell(3,1).value = 'Thang'
# ws.merge_cells('A3:H3') # Un-commenting this will corrupt!
# ws.merge_cells('A5:B5') # Un-commenting this will corrupt!

ws.insert_rows(7)
for col in range(1, 10):
    if ws.cell(8, col).has_style:
        ws.cell(7, col).font = copy(ws.cell(8, col).font)
        ws.cell(7, col).border = copy(ws.cell(8, col).border)

ws.cell(7, 1).value = 1
ws.cell(7, 2).value = 'Nguyen Van A'
ws.cell(7, 3).value = 1000
ws.cell(7, 4).value = '=IFERROR(C7*E7,"-")'

wb.save('my_custom_addons/dl_salary_kpi/scratch/test_corrupt_output.xlsx')
print("Saved simulated file.")
