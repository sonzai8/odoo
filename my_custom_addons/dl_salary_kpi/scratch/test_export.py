# -*- coding: utf-8 -*-
import io
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

def test_export():
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh Sach Nguoi Phu Thuoc"

    # Styles
    header_font = Font(bold=True)
    alignment = Alignment(horizontal='center', vertical='center')
    header_fill = openpyxl.styles.PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Header
    headers = [
        "STT", "Tên nhân viên", "Mã số thuế nhân viên", "Họ và tên người phụ thuộc", 
        "Quan hệ", "Mã số thuế NPT", "CCCD NPT", "Còn hiệu lực", "Ghi chú"
    ]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = header_font
        cell.alignment = alignment
        cell.fill = header_fill
        cell.border = border

    # Create Hidden Sheet
    ws_hidden = wb.create_sheet("DanhSachLoaiQuanHe")
    ws_hidden.sheet_state = 'hidden'
    relationships = ['Con', 'Vợ/Chồng', 'Bố/Mẹ', 'Anh/Chị/Em', 'Khác']
    for i, rel in enumerate(relationships, 1):
        ws_hidden.cell(row=i, column=1, value=rel)

    # Data Validation
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='=DanhSachLoaiQuanHe!$A$1:$A$5', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('E2:E1000')

    # Dummy Data
    for idx in range(1, 3):
        row = idx + 1
        data = [idx, "NV A", "123", "NPT A", "Con", "456", "789", "X", "Note"]
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical='center', horizontal='left' if col in [2, 4, 9] else 'center')

    # Table
    from openpyxl.worksheet.table import Table, TableStyleInfo
    last_row = 3
    table_range = f"A1:I{last_row}"
    table = Table(displayName="DanhSachNguoiPhuThuoc", ref=table_range)
    style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table.tableStyleInfo = style
    ws.add_table(table)

    wb.save(output)
    print("Export successful!")

if __name__ == "__main__":
    try:
        test_export()
    except Exception as e:
        print(f"Error: {e}")
