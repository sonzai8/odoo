import openpyxl

wb = openpyxl.load_workbook('my_custom_addons/dl_salary_kpi/static/src/templates/TEMPLATE_KPI_2026.xlsx')
ws = wb['TEMPLATE_KPI_POINT']

print("Merged cells in TEMPLATE_KPI_POINT:")
for m in ws.merged_cells.ranges:
    print(m)

print("\nRow 6 values:")
for col in range(1, 10):
    print(f"Col {col}: {ws.cell(row=6, column=col).value}")

print("\nRow 7 values:")
for col in range(1, 10):
    print(f"Col {col}: {ws.cell(row=7, column=col).value}")

print("\nRow 8 values:")
for col in range(1, 10):
    print(f"Col {col}: {ws.cell(row=8, column=col).value}")
