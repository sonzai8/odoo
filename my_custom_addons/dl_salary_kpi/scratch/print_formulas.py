from openpyxl import load_workbook
wb = load_workbook('test_corrupt_output.xlsx')
ws = wb.active
for col in range(1, 140):
    c8 = ws.cell(row=8, column=col)
    c9 = ws.cell(row=9, column=col)
    if c8.data_type == 'f':
        print(f"Col {col}:")
        print(f"  Row 8: {c8.value}")
        print(f"  Row 9: {c9.value}")
