# -*- coding: utf-8 -*-
from odoo import http
from odoo.http import request
import io
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, Protection, PatternFill
from .. import constants
from datetime import date
import calendar
from ..tools import no_accent_vietnamese

class AttendanceExportController(http.Controller):

    @http.route('/dl_salary_kpi/export_attendance/<int:month_id>', type='http', auth='user')
    def export_attendance(self, month_id, **kwargs):
        month = request.env['dl.salary.kpi.month'].browse(month_id)
        if not month.exists():
            return http.NotFound()

        output = io.BytesIO()
        wb = openpyxl.Workbook()
        wb.calculation.fullCalcOnLoad = True
        ws = wb.active
        ws.title = "Bang Cham Cong"
        ws.freeze_panes = 'D4'

        title_font = Font(size=16, bold=True)
        header_font = Font(bold=True)
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        alignment = Alignment(horizontal='center', vertical='center')
        
        sunday_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
        sunday_data_fill = PatternFill(start_color="E6F2FF", end_color="E6F2FF", fill_type="solid")
        fill_cp = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        fill_kp_o = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        fill_dc = PatternFill(start_color="E5CCFF", end_color="E5CCFF", fill_type="solid")
        fill_departed = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        no_fill = PatternFill(fill_type=None)

        ws.row_dimensions[1].height = 40
        
        # 1. Tiêu đề chính (A1:C1)
        ws.merge_cells('A1:C1')
        title_main = f"Chấm Công Thường Tháng {month.date_month.strftime('%m/%Y')}"
        cell_title = ws['A1']
        cell_title.value = title_main
        cell_title.font = Font(size=14, bold=True)
        cell_title.alignment = alignment
        
        # 2. Tên công ty (F1:AL1)
        ws.merge_cells('F1:AL1')
        cell_company = ws['F1']
        cell_company.value = month.company_id.name or "CÔNG TY ĐỨC LÂM"
        cell_company.font = Font(size=14, bold=True)
        cell_company.alignment = alignment

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 15

        headers_main = [
            constants.COL_STT, constants.COL_TAX_ID, constants.COL_FULL_NAME,
            constants.COL_BIRTHDAY, constants.COL_TAX_DEPARTMENT,
            constants.COL_POSITION, constants.COL_BASE_SALARY
        ]
        weekday_map = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]
        last_day = calendar.monthrange(month.date_month.year, month.date_month.month)[1]

        for col, text in enumerate(headers_main, 1):
            cell = ws.cell(row=2, column=col, value=text)
            cell.font = header_font
            cell.border = border
            cell.alignment = alignment
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

        for day in range(1, 32):
            col = 7 + day
            cell_day = ws.cell(row=2, column=col, value=f"{day:02d}")
            cell_day.font = header_font
            cell_day.border = border
            cell_day.alignment = alignment
            
            weekday_str = ""
            is_sun = False
            if day <= last_day:
                d = date(month.date_month.year, month.date_month.month, day)
                weekday_str = weekday_map[d.weekday()]
                is_sun = (d.weekday() == 6)
            
            cell_wd = ws.cell(row=3, column=col, value=weekday_str)
            cell_wd.font = header_font
            cell_wd.border = border
            cell_wd.alignment = alignment
            if is_sun:
                cell_day.fill = sunday_fill
                cell_wd.fill = sunday_fill

        row_num = 4
        for i, line in enumerate(month.line_ids.sorted(key=lambda l: (l.dl_tax_department_id.sequence or 999, l.employee_name or '')), 1):
            ws.row_dimensions[row_num].height = 30
            ws.cell(row=row_num, column=1, value=i).border = border
            ws.cell(row=row_num, column=2, value=line.dl_tax_id).border = border
            ws.cell(row=row_num, column=3, value=line.employee_name).border = border
            ws.cell(row=row_num, column=4, value=line.birthday).border = border
            ws.cell(row=row_num, column=5, value=line.dl_tax_department_id.name).border = border
            ws.cell(row=row_num, column=6, value=line.dl_tax_position).border = border
            ws.cell(row=row_num, column=7, value=line.dl_tax_base_salary).border = border
            
            departure_date = line.employee_id.dl_departure_date
            for day in range(1, 32):
                col_idx = 7 + day
                field_name = f"day_{day:02d}"
                att_type = getattr(line, field_name)
                code = att_type.code if att_type else ""
                cell = ws.cell(row=row_num, column=col_idx, value=code)
                cell.border = border
                cell.alignment = alignment
                
                d = None
                if day <= last_day:
                    d = date(month.date_month.year, month.date_month.month, day)
                
                if d and d.weekday() == 6:
                    cell.protection = Protection(locked=True)
                else:
                    cell.protection = Protection(locked=False)
                
                # Logic tô màu
                if departure_date and d and d >= departure_date:
                    cell.fill = fill_departed
                elif code == 'CP':
                    cell.fill = fill_cp
                elif code in ['KP', 'Ô']:
                    cell.fill = fill_kp_o
                elif code == 'ĐC':
                    cell.fill = fill_dc
                elif d and d.weekday() == 6:
                    cell.fill = sunday_data_fill
                else:
                    cell.fill = no_fill
                
            row = row_num
            # H:8, AL:38. AM:39, AN:40, AO:41, AP:42, AQ:43
            ws.cell(row=row, column=39, value=f'=COUNTIF(H{row}:AL{row},"N")+(COUNTIF(H{row}:AL{row},"N/1")+COUNTIF(H{row}:AL{row},"N/2"))*0.5').border = border
            ws.cell(row=row, column=40, value=f'=COUNTIF(H{row}:AL{row},"Đ")+(COUNTIF(H{row}:AL{row},"Đ/1")+COUNTIF(H{row}:AL{row},"Đ/2"))*0.5').border = border
            ws.cell(row=row, column=41, value=f'=AM{row}+AN{row}').border = border
            ws.cell(row=row, column=42, value=f'=COUNTIF(H{row}:AL{row},"PL")').border = border
            ws.cell(row=row, column=43, value=f'=COUNTIF(H{row}:AL{row},"P")').border = border
            row_num += 1

        summary_headers = [
            ("AM", constants.COL_ATT_NORMAL), ("AN", constants.COL_ATT_NIGHT),
            ("AO", constants.COL_ATT_TOTAL), ("AP", constants.COL_ATT_HOLIDAY),
            ("AQ", constants.COL_ATT_LEAVE)
        ]
        for i, (col_letter, text) in enumerate(summary_headers):
            col_idx = 39 + i
            cell = ws.cell(row=2, column=col_idx, value=text)
            cell.font = header_font
            cell.border = border
            cell.alignment = alignment
            ws.merge_cells(start_row=2, start_column=col_idx, end_row=3, end_column=col_idx)
            ws.column_dimensions[col_letter].width = 12

        ws_codes = wb.create_sheet("Ma cham cong")
        ws_codes.cell(row=1, column=1, value="Mã công").font = header_font
        ws_codes.cell(row=1, column=2, value="Tên loại công").font = header_font
        ws_codes.cell(row=1, column=3, value="Trọng số").font = header_font
        
        att_types = request.env['dl.salary.kpi.attendance.type'].search([('ot_type', '=', 'none')], order='weight desc')
        for idx, att in enumerate(att_types, 2):
            ws_codes.cell(row=idx, column=1, value=att.code)
            ws_codes.cell(row=idx, column=2, value=att.name)
            ws_codes.cell(row=idx, column=3, value=att.weight)
        
        from openpyxl.worksheet.datavalidation import DataValidation
        last_data_row = row_num - 1
        dv = DataValidation(type="list", formula1=f"'Ma cham cong'!$A$2:$A${len(att_types) + 1}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"H4:AL{last_data_row}")

        ws.auto_filter.ref = f"A3:AQ{last_data_row}"
        wb.save(output)
        output.seek(0)
        company_name = no_accent_vietnamese(month.company_id.name or "Duc_Lam")
        export_date = date.today().strftime('%d_%m_%Y')
        filename = f"Bang_Cham_Cong_{month.date_month.strftime('%m_%Y')}_{company_name}_{export_date}.xlsx"
        return request.make_response(output.getvalue(), headers=[('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), ('Content-Disposition', f'attachment; filename={filename}')])

    @http.route('/dl_salary_kpi/export_ot_attendance/<int:month_id>', type='http', auth='user')
    def export_ot_attendance(self, month_id, **kwargs):
        month = request.env['dl.salary.kpi.month'].browse(month_id)
        if not month.exists():
            return http.NotFound()

        output = io.BytesIO()
        wb = openpyxl.Workbook()
        wb.calculation.fullCalcOnLoad = True
        ws = wb.active
        ws.title = "Bang Cham Cong Lam Them"
        ws.freeze_panes = 'D4'

        title_font = Font(size=16, bold=True)
        header_font = Font(bold=True)
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        alignment = Alignment(horizontal='center', vertical='center')
        sunday_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
        sunday_data_fill = PatternFill(start_color="E6F2FF", end_color="E6F2FF", fill_type="solid")
        ot_header_fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        fill_cp = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        fill_kp_o = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        fill_dc = PatternFill(start_color="E5CCFF", end_color="E5CCFF", fill_type="solid")
        fill_departed = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        no_fill = PatternFill(fill_type=None)

        ws.row_dimensions[1].height = 40
        
        # 1. Tiêu đề chính (A1:C1)
        ws.merge_cells('A1:C1')
        title_main = f"Chấm Công Làm Thêm Tháng {month.date_month.strftime('%m/%Y')}"
        cell_title = ws['A1']
        cell_title.value = title_main
        cell_title.font = Font(size=14, bold=True)
        cell_title.alignment = alignment
        
        # 2. Tên công ty (F1:AL1)
        ws.merge_cells('F1:AL1')
        cell_company = ws['F1']
        cell_company.value = month.company_id.name or "CÔNG TY ĐỨC LÂM"
        cell_company.font = Font(size=14, bold=True)
        cell_company.alignment = alignment

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 15

        # Headers
        headers_main = [
            constants.COL_STT, constants.COL_TAX_ID, constants.COL_FULL_NAME,
            constants.COL_BIRTHDAY, constants.COL_TAX_DEPARTMENT,
            constants.COL_POSITION, constants.COL_BASE_SALARY
        ]
        for col, text in enumerate(headers_main, 1):
            cell = ws.cell(row=2, column=col, value=text)
            cell.font = header_font
            cell.border = border
            cell.alignment = alignment
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

        last_day = calendar.monthrange(month.date_month.year, month.date_month.month)[1]
        weekday_map = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]

        # Days Normal (8-38)
        for day in range(1, 32):
            col = 7 + day
            cell_day = ws.cell(row=2, column=col, value=f"{day:02d}")
            cell_day.font = header_font
            cell_day.border = border
            cell_day.alignment = alignment
            
            weekday_str = ""
            is_sun = False
            if day <= last_day:
                d = date(month.date_month.year, month.date_month.month, day)
                weekday_str = weekday_map[d.weekday()]
                is_sun = (d.weekday() == 6)
            
            cell_wd = ws.cell(row=3, column=col, value=weekday_str)
            cell_wd.font = header_font
            cell_wd.border = border
            cell_wd.alignment = alignment
            if is_sun:
                cell_day.fill = sunday_fill
                cell_wd.fill = sunday_fill

        # Summary Normal (39-43)
        summary_headers = [
            constants.COL_ATT_NORMAL, constants.COL_ATT_NIGHT, constants.COL_ATT_TOTAL,
            constants.COL_ATT_HOLIDAY, constants.COL_ATT_LEAVE
        ]
        for i, text in enumerate(summary_headers):
            col_idx = 39 + i
            cell = ws.cell(row=2, column=col_idx, value=text)
            cell.font = header_font
            cell.border = border
            cell.alignment = alignment
            ws.merge_cells(start_row=2, start_column=col_idx, end_row=3, end_column=col_idx)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 12

        # OT Days (44-74)
        for day in range(1, 32):
            col = 43 + day
            cell_day = ws.cell(row=2, column=col, value=f"{day:02d}")
            cell_day.font = header_font
            cell_day.border = border
            cell_day.alignment = alignment
            cell_day.fill = ot_header_fill
            
            weekday_str = ""
            is_sun = False
            if day <= last_day:
                d = date(month.date_month.year, month.date_month.month, day)
                weekday_str = weekday_map[d.weekday()]
                is_sun = (d.weekday() == 6)
            
            cell_wd = ws.cell(row=3, column=col, value=weekday_str)
            cell_wd.font = header_font
            cell_wd.border = border
            cell_wd.alignment = alignment
            cell_wd.fill = ot_header_fill
            if is_sun:
                cell_wd.font = Font(bold=True, color="FF0000")

        # OT Summary (75-77)
        ot_summary_headers = [constants.COL_OT_NORMAL, constants.COL_OT_NIGHT, constants.COL_OT_TOTAL]
        for i, text in enumerate(ot_summary_headers):
            col_idx = 75 + i
            cell = ws.cell(row=2, column=col_idx, value=text)
            cell.font = header_font
            cell.border = border
            cell.alignment = alignment
            ws.merge_cells(start_row=2, start_column=col_idx, end_row=3, end_column=col_idx)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

        row_num = 4
        for i, line in enumerate(month.line_ids.sorted(key=lambda l: (l.dl_tax_department_id.sequence or 999, l.employee_name or '')), 1):
            ws.row_dimensions[row_num].height = 30
            ws.cell(row=row_num, column=1, value=i).border = border
            ws.cell(row=row_num, column=2, value=line.dl_tax_id).border = border
            ws.cell(row=row_num, column=3, value=line.employee_name).border = border
            ws.cell(row=row_num, column=4, value=line.birthday).border = border
            ws.cell(row=row_num, column=5, value=line.dl_tax_department_id.name).border = border
            ws.cell(row=row_num, column=6, value=line.dl_tax_position).border = border
            ws.cell(row=row_num, column=7, value=line.dl_tax_base_salary).border = border
            
            departure_date = line.employee_id.dl_departure_date
            # Fill Normal Days (8-38)
            for day in range(1, 32):
                col_idx = 7 + day
                norm_att = getattr(line, f"day_{day:02d}")
                norm_code = norm_att.code if norm_att else ""
                cell = ws.cell(row=row_num, column=col_idx, value=norm_code)
                cell.border = border
                cell.alignment = alignment
                
                d = None
                if day <= last_day:
                    d = date(month.date_month.year, month.date_month.month, day)

                if departure_date and d and d >= departure_date:
                    cell.fill = fill_departed
                elif norm_code == 'CP':
                    cell.fill = fill_cp
                elif norm_code in ['KP', 'Ô']:
                    cell.fill = fill_kp_o
                elif norm_code == 'ĐC':
                    cell.fill = fill_dc
                elif d and d.weekday() == 6:
                    cell.fill = sunday_data_fill
                else:
                    cell.fill = no_fill

            row = row_num
            ws.cell(row=row, column=39, value=f'=COUNTIF(H{row}:AL{row},"N")+(COUNTIF(H{row}:AL{row},"N/1")+COUNTIF(H{row}:AL{row},"N/2"))*0.5').border = border
            ws.cell(row=row, column=40, value=f'=COUNTIF(H{row}:AL{row},"Đ")+(COUNTIF(H{row}:AL{row},"Đ/1")+COUNTIF(H{row}:AL{row},"Đ/2"))*0.5').border = border
            ws.cell(row=row, column=41, value=f'=AM{row}+AN{row}').border = border
            ws.cell(row=row, column=42, value=f'=COUNTIF(H{row}:AL{row},"PL")').border = border
            ws.cell(row=row, column=43, value=f'=COUNTIF(H{row}:AL{row},"P")').border = border

            # Fill OT Days (44-74)
            for day in range(1, 32):
                col_idx = 43 + day
                ot_att = getattr(line, f"ot_day_{day:02d}")
                ot_code = ot_att.code if ot_att else ""
                cell = ws.cell(row=row_num, column=col_idx, value=ot_code)
                cell.border = border
                cell.alignment = alignment
                
                d = None
                if day <= last_day:
                    d = date(month.date_month.year, month.date_month.month, day)
                
                if d and d.weekday() == 6:
                    cell.protection = Protection(locked=False)
                else:
                    cell.protection = Protection(locked=True)
                
                if departure_date and d and d >= departure_date:
                    cell.fill = fill_departed
                else:
                    cell.fill = ot_header_fill if not ot_att else sunday_data_fill

            # OT Formulas (H:8, AL:38... AR:44, BV:74)
            # BW:75, BX:76, BY:77
            formula_ot_n = f'=COUNTIF(AR{row}:BV{row},"0.5N")*0.5 + COUNTIF(AR{row}:BV{row},"CNN")*8 + COUNTIF(AR{row}:BV{row},"CNN/2")*4 + COUNTIF(AR{row}:BV{row},"LN")*8'
            ws.cell(row=row, column=75, value=formula_ot_n).border = border
            formula_ot_d = f'=COUNTIF(AR{row}:BV{row},"0.5Đ")*0.5 + COUNTIF(AR{row}:BV{row},"CNĐ")*8 + COUNTIF(AR{row}:BV{row},"CNĐ/2")*4 + COUNTIF(AR{row}:BV{row},"LĐ")*8'
            ws.cell(row=row, column=76, value=formula_ot_d).border = border
            ws.cell(row=row, column=77, value=f'=BW{row}+BX{row}').border = border
            row_num += 1
            
        last_data_row = row_num - 1
        ws.auto_filter.ref = f"A3:BY{last_data_row}"

        ws_codes_ot = wb.create_sheet("Ma tang ca")
        ws_codes_ot.cell(row=1, column=1, value="Mã tăng ca").font = header_font
        ws_codes_ot.cell(row=1, column=2, value="Tên loại").font = header_font
        att_types_ot = request.env['dl.salary.kpi.attendance.type'].search([('ot_type', '!=', 'none')], order='weight desc')
        for idx, att in enumerate(att_types_ot, 2):
            ws_codes_ot.cell(row=idx, column=1, value=att.code)
            ws_codes_ot.cell(row=idx, column=2, value=att.name)

        from openpyxl.worksheet.datavalidation import DataValidation
        dv_ot = DataValidation(type="list", formula1=f"'Ma tang ca'!$A$2:$A${len(att_types_ot) + 1}", allow_blank=True)
        ws.add_data_validation(dv_ot)
        dv_ot.add(f"H4:AL{last_data_row}")
        dv_ot.add(f"AR4:BV{last_data_row}")

        wb.save(output)
        output.seek(0)
        company_name = no_accent_vietnamese(month.company_id.name or "Duc_Lam")
        export_date = date.today().strftime('%d_%m_%Y')
        filename = f"Cong_Lam_Them_{month.date_month.strftime('%m_%Y')}_{company_name}_{export_date}.xlsx"
        return request.make_response(output.getvalue(), headers=[('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), ('Content-Disposition', f'attachment; filename={filename}')])
