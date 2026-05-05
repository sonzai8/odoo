# -*- coding: utf-8 -*-
from odoo import http
from odoo.http import request
import io
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, Protection
from .. import constants
import logging

_logger = logging.getLogger(__name__)

class EmployeeExportController(http.Controller):

    @http.route('/dl_salary_kpi/export_tax_employees', type='http', auth='user')
    def export_tax_employees(self, **kwargs):
        employees = request.env['hr.employee'].search([('active', 'in', [True, False])], order='dl_tax_department_id, dl_first_name')
        
        output = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh Sach Nhan Vien Thue"

        header_font = Font(bold=True)
        alignment = Alignment(horizontal='center', vertical='center')
        header_fill = openpyxl.styles.PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        headers = [
            constants.COL_STT, constants.COL_ID_NV, constants.COL_FULL_NAME, constants.COL_FIRST_NAME,
            constants.COL_CCCD, constants.COL_BIRTHDAY, constants.COL_EMAIL, constants.COL_PHONE,
            constants.COL_GENDER, constants.COL_TAX_ID, constants.COL_TAX_POSITION,
            constants.COL_TAX_DEPARTMENT, constants.COL_TAX_BASE_SALARY, constants.COL_DEPARTURE_DATE
        ]
        ws.row_dimensions[1].height = 30
        for col, text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=text)
            cell.font = header_font
            cell.alignment = alignment
            cell.fill = header_fill
            cell.border = border

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 5
        ws.column_dimensions['H'].width = 15
        ws.column_dimensions['I'].width = 12
        ws.column_dimensions['J'].width = 18
        ws.column_dimensions['K'].width = 18
        ws.column_dimensions['L'].width = 30
        ws.column_dimensions['M'].width = 20
        ws.column_dimensions['N'].width = 15

        for idx, emp in enumerate(employees, 1):
            row = idx + 1
            ws.row_dimensions[row].height = 25
            data = [idx, emp.id, emp.name, emp.dl_first_name, emp.identification_id, emp.birthday, emp.email, emp.work_phone, 'Nam' if emp.sex == 'male' else 'Nữ' if emp.sex == 'female' else 'Khác', emp.dl_tax_id, emp.dl_tax_position, emp.dl_tax_department_id.name if emp.dl_tax_department_id else '', emp.dl_tax_base_salary, emp.dl_departure_date]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical='center', horizontal='left' if col in [3, 5] else 'center')
                if col in [7, 8, 10]:
                    cell.number_format = '@'
                    if value: cell.value = str(value)

        from openpyxl.worksheet.table import Table, TableStyleInfo
        last_row = len(employees) + 1
        if last_row > 1:
            table = Table(displayName="DanhSachNhanVienThue", ref=f"A1:N{last_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)

        wb.save(output)
        output.seek(0)
        filename = "Danh_Sach_Nhan_Vien_Thue.xlsx"
        return request.make_response(output.getvalue(), headers=[('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), ('Content-Disposition', f'attachment; filename={filename}')])

    @http.route('/dl_salary_kpi/export_dependents', type='http', auth='user')
    def export_dependents(self, **kwargs):
        try:
            dependents = request.env['dl.dependent'].search([], order='employee_id, name')
            output = io.BytesIO()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Danh Sach Nguoi Phu Thuoc"

            headers = [
                constants.COL_STT, constants.COL_EMP_NAME, constants.COL_EMP_TAX_ID,
                constants.COL_DEP_NAME, constants.COL_RELATIONSHIP, constants.COL_DEP_TAX_ID,
                constants.COL_DEP_CCCD, constants.COL_IS_ACTIVE, constants.COL_NOTE
            ]
            header_font = Font(bold=True)
            alignment = Alignment(horizontal='center', vertical='center')
            header_fill = openpyxl.styles.PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
            border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

            ws.row_dimensions[1].height = 30
            for col, text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=text)
                cell.font = header_font
                cell.alignment = alignment
                cell.fill = header_fill
                cell.border = border

            ws.column_dimensions['A'].width = 6
            ws.column_dimensions['B'].width = 25
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 25
            ws.column_dimensions['E'].width = 15
            ws.column_dimensions['F'].width = 20
            ws.column_dimensions['G'].width = 20
            ws.column_dimensions['H'].width = 15
            ws.column_dimensions['I'].width = 30

            relationship_map = {'child': 'Con', 'spouse': 'Vợ/Chồng', 'parent': 'Bố/Mẹ', 'sibling': 'Anh/Chị/Em', 'other': 'Khác'}
            
            ws_hidden = wb.create_sheet("DanhSachLoaiQuanHe")
            ws_hidden.sheet_state = 'hidden'
            relationships = ['Con', 'Vợ/Chồng', 'Bố/Mẹ', 'Anh/Chị/Em', 'Khác']
            for i, rel in enumerate(relationships, 1):
                ws_hidden.cell(row=i, column=1, value=rel)

            from openpyxl.worksheet.datavalidation import DataValidation
            dv = DataValidation(type="list", formula1='=DanhSachLoaiQuanHe!$A$1:$A$5', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add('E2:E1000')

            for idx, dep in enumerate(dependents, 1):
                row = idx + 1
                ws.row_dimensions[row].height = 25
                data = [idx, dep.employee_id.name or '', dep.employee_id.dl_tax_id or '', dep.name or '', relationship_map.get(dep.relationship, dep.relationship or ''), dep.dependent_number or '', dep.dependent_id_card or '', 'X' if dep.active else '', dep.note or '']
                for col, value in enumerate(data, 1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.border = border
                    cell.alignment = Alignment(vertical='center', horizontal='left' if col in [2, 4, 9] else 'center')
                    if col in [3, 6, 7]: cell.number_format = '@'

            from openpyxl.worksheet.table import Table, TableStyleInfo
            last_row = len(dependents) + 1
            if last_row > 1:
                table = Table(displayName="DanhSachNguoiPhuThuoc", ref=f"A1:I{last_row}")
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(table)

            wb.save(output)
            output.seek(0)
            filename = "Danh_Sach_Nguoi_Phu_Thuoc.xlsx"
            return request.make_response(output.getvalue(), headers=[('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), ('Content-Disposition', f'attachment; filename={filename}')])
        except Exception as e:
            _logger.error(f"Error exporting dependents: {e}", exc_info=True)
            return request.make_response(f"Error: {str(e)}", status=500)
