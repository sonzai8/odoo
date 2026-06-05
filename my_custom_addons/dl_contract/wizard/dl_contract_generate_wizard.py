# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import io
import base64
from datetime import date

try:
    from docx import Document
except ImportError:
    Document = None

class DlContractGenerateWizard(models.TransientModel):
    _name = 'dl.contract.generate.wizard'
    _description = 'Tạo File Hợp đồng từ Template'

    contract_id = fields.Many2one('dl.contract', string='Hợp đồng', required=True)
    template_id = fields.Many2one('dl.contract.template', string='Template HĐ', required=True, domain="['|', ('company_id', '=', False), ('company_id', '=', company_id)]")
    company_id = fields.Many2one('res.company', string='Công ty', related='contract_id.company_id')
    
    file_data = fields.Binary(string='File tải về', readonly=True)
    file_name = fields.Char(string='Tên file', readonly=True)
    state = fields.Selection([
        ('choose', 'Chọn Template'),
        ('done', 'Đã tạo xong')
    ], default='choose')

    @api.model
    def default_get(self, fields_list):
        res = super(DlContractGenerateWizard, self).default_get(fields_list)
        contract_id = self._context.get('default_contract_id') or res.get('contract_id')
        if contract_id:
            contract = self.env['dl.contract'].browse(contract_id)
            if 'template_id' in fields_list and not res.get('template_id'):
                if contract.contract_type_id and contract.contract_type_id.template_id:
                    res['template_id'] = contract.contract_type_id.template_id.id
        return res

    def action_generate(self):
        self.ensure_one()
        if not Document:
            raise UserError(_('Thư viện python-docx chưa được cài đặt. Vui lòng liên hệ Admin.'))
        
        if not self.template_id.template_file:
            raise UserError(_('Template này chưa có file mẫu. Vui lòng upload file mẫu (.docx) trong cấu hình template.'))

        # Decode template file
        template_bytes = base64.b64decode(self.template_id.template_file)
        doc_stream = io.BytesIO(template_bytes)
        
        try:
            doc = Document(doc_stream)
        except Exception as e:
            raise UserError(_('Lỗi khi đọc file mẫu. Đảm bảo file là định dạng .docx hợp lệ. Chi tiết lỗi: %s') % str(e))

        # Chuẩn bị dữ liệu thay thế
        contract = self.contract_id
        employee = contract.employee_id
        company = contract.company_id
        
        # Tính thời hạn hợp đồng từ loại hợp đồng default_duration_months
        contract_type = contract.contract_type_id
        if contract_type.duration_type == 'indefinite':
            term_str = 'Không xác định thời hạn'
        else:
            term_str = f"{contract_type.default_duration_months or 12} tháng"

        # Tính dạng text của ngày ký hợp đồng (ngày bắt đầu hiệu lực)
        text_contract_date_str = ''
        if contract.date_start:
            text_contract_date_str = f"ngày {contract.date_start.strftime('%d')} tháng {contract.date_start.strftime('%m')} năm {contract.date_start.strftime('%Y')}"

        # Lấy nhãn hiển thị của Nơi cấp CCCD (Selection)
        cccd_place_label = ''
        if getattr(employee, 'x_cccd_place', False):
            cccd_place_label = dict(employee._fields['x_cccd_place']._description_selection(self.env)).get(employee.x_cccd_place, '')

        replace_dict = {
            '{{employee_name}}': employee.name or '',
            '{{employee_cccd}}': employee.identification_id or '',
            '{{employee_cccd_date}}': employee.x_cccd_date.strftime('%d/%m/%Y') if getattr(employee, 'x_cccd_date', False) and employee.x_cccd_date else '',
            '{{employee_cccd_place}}': cccd_place_label,
            '{{employee_birthday}}': employee.birthday.strftime('%d/%m/%Y') if employee.birthday else '',
            '{{employee_gender}}': 'Nam' if employee.sex == 'male' else 'Nữ' if employee.sex == 'female' else '',
            '{{employee_address}}': employee.private_street or '', # Cần check trường địa chỉ thực tế của hr.employee
            '{{employee_phone}}': employee.private_phone or employee.work_phone or employee.mobile_phone or '',
            '{{employee_dept}}': contract.department_id.name or employee.department_id.name or '',
            '{{employee_salary}}': '{:,.0f}'.format(contract.total_wage).replace(',', '.') or '0',
            '{{employee_job_title}}': contract.job_title or getattr(employee, 'job_title', '') or '',
            '{{contract_number}}': contract.name or '',
            '{{contract_type}}': contract_type.name or '',
            '{{contract_term}}': term_str,
            '{{contract_start_date}}': contract.date_start.strftime('%d/%m/%Y') if contract.date_start else '',
            '{{contract_end_date}}': contract.date_end.strftime('%d/%m/%Y') if contract.date_end else '...',
            '{{text_contract_date}}': text_contract_date_str,
            '{{date_start}}': contract.date_start.strftime('%d/%m/%Y') if contract.date_start else '',
            '{{date_end}}': contract.date_end.strftime('%d/%m/%Y') if contract.date_end else '...',
            '{{job_title}}': contract.job_title or '',
            '{{department}}': contract.department_id.name or '',
            '{{wage}}': '{:,.0f}'.format(contract.wage).replace(',', '.') or '0',
            '{{allowance}}': '{:,.0f}'.format(contract.allowance).replace(',', '.') or '0',
            '{{total_wage}}': '{:,.0f}'.format(contract.total_wage).replace(',', '.') or '0',
            '{{company_name}}': company.name or '',
            '{{company_address}}': company.street or '',
            '{{today}}': fields.Date.today().strftime('%d/%m/%Y'),
        }

        # Hàm thay thế văn bản trong các run (để giữ format)
        def replace_text_in_paragraph(paragraph, replace_dict):
            text = paragraph.text
            for k, v in replace_dict.items():
                if k in text:
                    # Đơn giản hóa: replace thẳng trên paragraph.text (làm mất format chi tiết nếu có trong 1 từ)
                    # Nếu cần giữ format cao cấp, cần duyệt từng run và check cẩn thận. Ở đây tạm dùng cách thay text
                    pass
            # Cách an toàn hơn để giữ format cơ bản:
            for k, v in replace_dict.items():
                if k in paragraph.text:
                    inline = paragraph.runs
                    for i in range(len(inline)):
                        if k in inline[i].text:
                            text = inline[i].text.replace(k, str(v))
                            inline[i].text = text
            # Fallback nếu placeholder bị cắt ra nhiều run
            for k, v in replace_dict.items():
                if k in paragraph.text:
                    paragraph.text = paragraph.text.replace(k, str(v))

        # Duyệt qua các paragraph
        for p in doc.paragraphs:
            replace_text_in_paragraph(p, replace_dict)

        # Duyệt qua các table
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        replace_text_in_paragraph(p, replace_dict)

        # Lưu file
        out_stream = io.BytesIO()
        doc.save(out_stream)
        out_stream.seek(0)
        
        file_content = base64.b64encode(out_stream.read())
        safe_filename = f"HDLD_{employee.name.replace(' ', '_')}_{fields.Date.today().strftime('%Y%m%d')}.docx"

        self.write({
            'file_data': file_content,
            'file_name': safe_filename,
            'state': 'done'
        })
        
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'dl.contract.generate.wizard',
            'view_mode': 'form',
            'res_id': self.id,
            'views': [(False, 'form')],
            'target': 'new',
        }
