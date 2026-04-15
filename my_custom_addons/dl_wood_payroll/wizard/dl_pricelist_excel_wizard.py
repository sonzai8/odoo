# -*- coding: utf-8 -*-
import base64
import io
import openpyxl
from odoo import models, fields, api, _
from odoo.exceptions import UserError

class PricelistExcelWizard(models.TransientModel):
    _name = 'dl.pricelist.excel.wizard'
    _description = 'Wizard Nhập/Xuất Excel Bảng giá'

    file_data = fields.Binary(string='File Excel')
    file_name = fields.Char(string='Tên file')
    pricelist_id = fields.Many2one('dl.piece.rate.pricelist', string='Bảng giá', required=True)
    
    import_mode = fields.Selection([
        ('flat', 'Danh sách phẳng (Mặc định)'),
        ('matrix', 'Ma trận (Công đoạn là Cột)'),
    ], string='Định dạng Excel', default='flat', required=True)

    state = fields.Selection([
        ('draft', 'Nháp'),
        ('verified', 'Đã kiểm tra'),
    ], string='Trạng thái', default='draft')

    line_ids = fields.One2many('dl.pricelist.excel.line', 'wizard_id', string='Dòng xem trước')

    def action_export(self):
        """Xuất dữ liệu bảng giá hiện tại ra file Excel"""
        self.ensure_one()
        if self.import_mode == 'matrix':
            return self.action_export_matrix()
            
        # Tạo workbook mới (Flat List)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Don_Gia_Phang"
        # ... existing export logic ...
        headers = [
            'ID Dòng', 'Công đoạn (Bộ phận)', 'Mã sản phẩm (Internal Ref)', 
            'Tên sản phẩm', 'Giá thấp (Thiếu công)', 'Giá cao (Đủ công)', 'Giá lũy tiến'
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

        row = 2
        for line in self.pricelist_id.line_ids:
            ws.cell(row=row, column=1, value=line.id)
            ws.cell(row=row, column=2, value=line.department_id.name)
            ws.cell(row=row, column=3, value=line.product_id.default_code or '')
            ws.cell(row=row, column=4, value=line.product_id.name)
            ws.cell(row=row, column=5, value=line.price_low)
            ws.cell(row=row, column=6, value=line.price_high)
            ws.cell(row=row, column=7, value=line.extra_price)
            row += 1

        return self._save_and_download_wb(wb, f"Bang_gia_phang_{self.pricelist_id.name}.xlsx")

    def action_export_matrix(self):
        """Xuất bảng giá dưới dạng ma trận: SP là dòng, Công đoạn là cột"""
        self.ensure_one()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ma_Tran_Don_Gia"

        # Lấy danh sách công đoạn có trong hệ thống hoặc bảng giá
        departments = self.env['hr.department'].search([('x_is_production_stage', '=', True)], order='name')
        
        # Header cố định
        ws.cell(row=1, column=1, value='Mã sản phẩm').font = openpyxl.styles.Font(bold=True)
        ws.cell(row=1, column=2, value='Tên sản phẩm').font = openpyxl.styles.Font(bold=True)
        
        # Header công đoạn (Cột động)
        col = 3
        dept_col_map = {}
        for dept in departments:
            # Shift 3 cột cho mỗi công đoạn: Thấp, Cao, Lũy tiến
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+2)
            cell = ws.cell(row=1, column=col, value=dept.name)
            cell.font = openpyxl.styles.Font(bold=True)
            cell.alignment = openpyxl.styles.Alignment(horizontal='center')
            
            ws.cell(row=2, column=col, value=f"{dept.name} [Low]")
            ws.cell(row=2, column=col+1, value=f"{dept.name} [High]")
            ws.cell(row=2, column=col+2, value=f"{dept.name} [Extra]")
            
            dept_col_map[dept.id] = col
            col += 3

        # Lấy danh sách sản phẩm từ bảng giá hiện tại
        products = self.pricelist_id.line_ids.mapped('product_id')
        
        row = 3
        for prod in products:
            ws.cell(row=row, column=1, value=prod.default_code or '')
            ws.cell(row=row, column=2, value=prod.name)
            
            # Đổ giá cho từng công đoạn
            prod_lines = self.pricelist_id.line_ids.filtered(lambda l: l.product_id.id == prod.id)
            for line in prod_lines:
                if line.department_id.id in dept_col_map:
                    c = dept_col_map[line.department_id.id]
                    ws.cell(row=row, column=c, value=line.price_low)
                    ws.cell(row=row, column=c+1, value=line.price_high)
                    ws.cell(row=row, column=c+2, value=line.extra_price)
            row += 1

        return self._save_and_download_wb(wb, f"Ma_tran_gia_{self.pricelist_id.name}.xlsx")

    def _save_and_download_wb(self, wb, file_name):
        fp = io.BytesIO()
        wb.save(fp)
        data = base64.b64encode(fp.getvalue())
        fp.close()
        self.write({'file_data': data, 'file_name': file_name})
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model=dl.pricelist.excel.wizard&id={self.id}&field=file_data&download=true&filename={self.file_name}',
            'target': 'self',
        }

    def action_verify(self):
        """Kiểm tra dữ liệu (Tự động nhận diện Ma trận hoặc Phẳng)"""
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Vui lòng chọn file Excel."))

        try:
            fp = io.BytesIO(base64.b64decode(self.file_data))
            wb = openpyxl.load_workbook(fp, data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError(_("Lỗi đọc file: %s") % str(e))

        self.line_ids.unlink()
        
        # Nhận diện mode dựa trên header
        first_row = [cell.value for cell in ws[1]]
        if 'ID Dòng' in first_row or 'Công đoạn (Bộ phận)' in first_row:
            return self._verify_flat(ws)
        else:
            return self._verify_matrix(ws)

    def _verify_flat(self, ws):
        preview_lines = []
        product_cache = {}
        dept_cache = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row): continue
            line_id, dept_name, prod_ref, prod_name, p_low, p_high, p_extra = row[0:7]
            status, msg, dept, prod = self._check_line(dept_name, prod_ref or prod_name, p_low, p_high, dept_cache, product_cache)
            preview_lines.append((0, 0, {
                'line_id': int(line_id) if line_id else False,
                'dept_name': dept_name, 'product_ref': prod_ref, 'product_name': prod_name,
                'price_low': float(p_low or 0), 'price_high': float(p_high or 0), 'extra_price': float(p_extra or 0),
                'status': status, 'message': msg
            }))
        return self._finalize_verification(preview_lines)

    def _verify_matrix(self, ws):
        """Verify logic cho định dạng ma trận"""
        preview_lines = []
        product_cache = {}
        dept_cache = {}
        
        # Header hàng 1: Công đoạn
        # Header hàng 2: Low/High/Extra (nếu có)
        headers_1 = [cell.value for cell in ws[1]]
        headers_2 = [cell.value for cell in ws[2]]
        
        # Tìm các cột công đoạn
        dept_cols = []
        for i, h in enumerate(headers_1):
            if i < 2 or not h: continue # Bỏ qua Mã SP, Tên SP
            # h là tên công đoạn
            dept = self.env['hr.department'].search([('name', '=', str(h).strip()), ('x_is_production_stage', '=', True)], limit=1)
            if dept:
                dept_cols.append({'dept_id': dept.id, 'dept_name': dept.name, 'col_idx': i})

        for row in ws.iter_rows(min_row=3, values_only=True):
            if not any(row): continue
            prod_ref = str(row[0] or '').strip()
            prod_name = str(row[1] or '').strip()
            
            for dc in dept_cols:
                idx = dc['col_idx']
                p_low = row[idx]
                p_high = row[idx+1] if idx+1 < len(row) else 0.0
                p_extra = row[idx+2] if idx+2 < len(row) else 0.0
                
                if p_low or p_high or p_extra:
                    status, msg, dept, prod = self._check_line(dc['dept_name'], prod_ref or prod_name, p_low, p_high, dept_cache, product_cache)
                    preview_lines.append((0, 0, {
                        'dept_name': dc['dept_name'],
                        'product_ref': prod_ref,
                        'product_name': prod_name,
                        'price_low': float(p_low or 0),
                        'price_high': float(p_high or 0),
                        'extra_price': float(p_extra or 0),
                        'status': status,
                        'message': msg
                    }))
        return self._finalize_verification(preview_lines)

    def _check_line(self, dept_name, prod_key, p_low, p_high, dept_cache, prod_cache):
        status = 'ready'
        msgs = []
        dept = False
        prod = False
        
        if dept_name in dept_cache: dept = dept_cache[dept_name]
        else:
            dept = self.env['hr.department'].search([('name', '=', dept_name)], limit=1)
            dept_cache[dept_name] = dept
        
        if prod_key in prod_cache: prod = prod_cache[prod_key]
        else:
            prod = self.env['product.product'].search(['|', ('default_code', '=', prod_key), ('name', '=', prod_key)], limit=1)
            prod_cache[prod_key] = prod
            
        if not dept:
            status = 'error'; msgs.append(_("Thiếu công đoạn '%s'") % dept_name)
        if not prod:
            status = 'error'; msgs.append(_("Thiếu sản phẩm '%s'") % prod_key)
        if float(p_low or 0) < 0 or float(p_high or 0) < 0:
            status = 'error'; msgs.append(_("Giá không được âm"))
            
        return status, ". ".join(msgs) if msgs else _("OK"), dept, prod

    def _finalize_verification(self, preview_lines):
        self.write({'line_ids': preview_lines, 'state': 'verified'})
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'dl.pricelist.excel.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_import(self):
        """Import logic tương tự, dùng chung cho cả 2 mode vì preview lines là dạng phẳng"""
        self.ensure_one()
        # ... logic import từ line_ids ...
        if self.state != 'verified': return self.action_verify()
        
        valid_lines = self.line_ids.filtered(lambda l: l.status != 'error')
        success, update = 0, 0
        for pl in valid_lines:
            dept = self.env['hr.department'].search([('name', '=', pl.dept_name)], limit=1)
            prod = self.env['product.product'].search(['|', ('default_code', '=', pl.product_ref), ('name', '=', pl.product_ref or pl.product_name)], limit=1)
            if not dept or not prod: continue
            
            vals = {'pricelist_id': self.pricelist_id.id, 'department_id': dept.id, 'product_id': prod.id,
                    'price_low': pl.price_low, 'price_high': pl.price_high, 'extra_price': pl.extra_price}
            
            # Update by ID or Duplicate search
            existing = False
            if pl.line_id:
                existing = self.env['dl.piece.rate.pricelist.line'].browse(pl.line_id)
                if existing.exists() and existing.pricelist_id.id == self.pricelist_id.id:
                    existing.write(vals); update += 1; existing = True
                else: existing = False
            
            if not existing:
                dup = self.env['dl.piece.rate.pricelist.line'].search([
                    ('pricelist_id', '=', self.pricelist_id.id), ('department_id', '=', dept.id), ('product_id', '=', prod.id)
                ], limit=1)
                if dup: dup.write(vals); update += 1
                else: self.env['dl.piece.rate.pricelist.line'].create(vals); success += 1
        
        return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {
            'title': _('Thành công'), 'message': _('Đã tạo %d và cập nhật %d dòng.') % (success, update), 'type': 'success'}}

class PricelistExcelLine(models.TransientModel):
    _name = 'dl.pricelist.excel.line'
    _description = 'Dòng xem trước từ Excel'

    wizard_id = fields.Many2one('dl.pricelist.excel.wizard', ondelete='cascade')
    line_id = fields.Integer(string='ID Dòng Gốc')
    
    dept_name = fields.Char(string='Công đoạn')
    product_ref = fields.Char(string='Mã sản phẩm')
    product_name = fields.Char(string='Tên sản phẩm')
    
    price_low = fields.Float(string='Giá thấp')
    price_high = fields.Float(string='Giá cao')
    extra_price = fields.Float(string='Giá lũy tiến')
    
    status = fields.Selection([
        ('ready', 'Hợp lệ'),
        ('warning', 'Cảnh báo'),
        ('error', 'Lỗi'),
    ], string='Trạng thái', default='ready')
    message = fields.Text(string='Thông báo')
